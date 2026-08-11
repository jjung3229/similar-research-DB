#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""수용할 또는 사용할 토지세목조서(PDF) → 엑셀 변환기.

국토교통부·지자체 고시문에 붙는 「수용할 또는 사용할 토지세목조서」 표를
PDF에서 그대로 읽어 엑셀(.xlsx)로 옮긴다.

- 표의 머리글(번호/소재지/지번/지목/지적/편입면적/소유자/이해관계인/비고)을
  읽어서 열 구성을 자동으로 잡는다. 고시문마다 열이 조금씩 달라도 따라간다.
- 한 건이 여러 줄에 걸쳐 인쇄된 것(소재지가 2줄, 주소가 3줄 등)을 한 행으로 합친다.
- 페이지마다 반복되는 머리글은 버리고, '소 계' 행은 따로 모아둔다.
- 마지막에 각 소계 금액과 직접 합산한 값을 비교해서 누락 여부를 검증한다.

사용법:
    python land_schedule.py 고시문.pdf                 # 같은 이름의 .xlsx 생성
    python land_schedule.py 고시문.pdf -o 결과.xlsx
    python land_schedule.py *.pdf -d out/             # 여러 개 한꺼번에
"""

from __future__ import annotations

import argparse
import glob
import os
import re
import sys
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence, Tuple

try:
    import pdfplumber
except ImportError:  # pragma: no cover
    sys.exit("pdfplumber 가 필요합니다.  pip install -r requirements.txt")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl 이 필요합니다.  pip install -r requirements.txt")


# ---------------------------------------------------------------------------
# 표 머리글 해석
# ---------------------------------------------------------------------------

#: 번호 열에 들어오면 자료 행이 아니라 합계 행으로 본다.
SUBTOTAL_LABELS = ("소계", "소 계", "합계", "합 계", "총계", "총 계", "계")

#: 머리글이 이 낱말들을 담고 있으면 토지세목조서 표로 판단한다.
REQUIRED_HEADER_WORDS = ("번호", "지번")

#: 줄바꿈을 공백 없이 이어붙일 열(사람·법인 이름처럼 한 낱말이 접힌 경우).
GLUE_WITHOUT_SPACE = ("성명", "지목", "지번", "권리")

#: 숫자로 바꿔 쓸 열.
NUMERIC_HINTS = ("면적", "지적", "㎡", "m2")


def _norm(text: Optional[str]) -> str:
    """셀 문자열 정리: 공백/줄바꿈 정돈."""
    if text is None:
        return ""
    return re.sub(r"[ \t]+", " ", text.replace("\xa0", " ")).strip()


def _flatten(text: str, glue: str) -> str:
    """셀 안의 줄바꿈을 지정한 방식으로 이어붙인다."""
    parts = [p.strip() for p in text.split("\n") if p.strip()]
    return glue.join(parts)


@dataclass
class TableLayout:
    """표의 열 구성. 머리글 행에서 만들어진다."""

    columns: List[str]              # 최종 열 이름 (예: '소유자 성명')
    col_index: List[int]            # 원본 표에서의 열 번호
    key_col: int                    # '번호' 열이 columns 안에서 갖는 위치
    header_rows: int                # 머리글 앞뒤로 버릴 줄 수
    anchor_col: Optional[int] = None  # 항목마다 반드시 값이 있는 열('지번')

    def signature(self) -> Tuple[str, ...]:
        return tuple(self.columns)


def _header_span(rows: Sequence[Sequence[Optional[str]]]) -> int:
    """머리글이 몇 줄인지 판단한다(보통 2줄)."""
    span = 1
    if len(rows) > 1:
        second = [_norm(c) for c in rows[1]]
        # 두 번째 줄이 '시군읍면동리 / 성명 / 주소' 같은 하위 머리글이면 머리글에 포함
        if any(v in ("시군읍면동리", "성명", "주소", "권리의 종류", "권리의종류") for v in second):
            span = 2
    return span


def _build_layout(rows: Sequence[Sequence[Optional[str]]]) -> Optional[TableLayout]:
    """표의 첫 줄(들)을 읽어 열 구성을 만든다."""
    if not rows:
        return None

    span = _header_span(rows)
    head = [[_norm(c) for c in row] for row in rows[:span]]
    width = max(len(r) for r in head)
    for r in head:
        r.extend([""] * (width - len(r)))

    top_flat = " ".join(head[0])
    if not all(word in top_flat for word in REQUIRED_HEADER_WORDS):
        return None

    # '소 유 자'처럼 두 칸을 아우르는 머리글은 병합 셀이라 오른쪽 칸이 비어 있다.
    # 하위 머리글('성명'/'주소')에 붙여 주려고 앞의 값을 끌고 온다.
    top: List[str] = []
    carry = ""
    for value in head[0]:
        if value:
            carry = value
        top.append(carry)

    columns: List[str] = []
    col_index: List[int] = []
    for idx in range(width):
        upper = _flatten(head[0][idx], "")
        lower = _flatten(head[1][idx], "") if span > 1 else ""
        merged_upper = _flatten(top[idx], "")

        if lower and merged_upper and lower != merged_upper:
            name = f"{merged_upper.replace(' ', '')} {lower}"
        elif lower:
            name = lower
        elif upper:
            name = upper
        else:
            # 병합 셀의 이어지는 칸이거나 표 바깥의 여백 열 → 버린다.
            name = ""

        if not name:
            continue
        name = name.replace("\n", " ").strip()
        columns.append(name)
        col_index.append(idx)

    # 같은 이름이 두 번 나오면 뒤에 번호를 붙인다.
    seen: Dict[str, int] = {}
    for i, name in enumerate(columns):
        if name in seen:
            seen[name] += 1
            columns[i] = f"{name}{seen[name]}"
        else:
            seen[name] = 1

    try:
        key_col = next(i for i, c in enumerate(columns) if c.startswith("번호"))
    except StopIteration:
        return None

    anchor_col = next((i for i, c in enumerate(columns) if c.startswith("지번")), None)

    return TableLayout(
        columns=columns,
        col_index=col_index,
        key_col=key_col,
        header_rows=span,
        anchor_col=anchor_col,
    )


#: 표 위쪽에 고시 본문이 함께 잡히는 경우가 있어 머리글을 몇 줄까지 찾아볼지.
MAX_HEADER_SEARCH = 6


def locate_layout(table: Sequence[Sequence[Optional[str]]]) -> Optional[TableLayout]:
    """표 안에서 머리글 줄을 찾아 열 구성을 만든다.

    첫 페이지처럼 고시 본문이 표 첫 줄에 딸려 들어오는 경우가 있어
    맨 위 몇 줄을 건너뛰며 머리글을 찾는다. header_rows 는 '본문 시작 전까지
    버릴 줄 수'가 된다.
    """
    for offset in range(min(MAX_HEADER_SEARCH, len(table))):
        layout = _build_layout(table[offset:])
        if layout is not None:
            layout.header_rows += offset
            return layout
    return None


# ---------------------------------------------------------------------------
# 본문 추출
# ---------------------------------------------------------------------------


@dataclass
class ExtractResult:
    columns: List[str] = field(default_factory=list)
    records: List[Dict[str, Any]] = field(default_factory=list)
    subtotals: List[Dict[str, Any]] = field(default_factory=list)
    #: ('record'|'subtotal', 행) 을 원문에 인쇄된 순서 그대로 담는다. 소계 검증에 쓴다.
    sequence: List[Tuple[str, Dict[str, Any]]] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    pages_with_table: int = 0
    total_pages: int = 0


TABLE_SETTINGS = {
    "vertical_strategy": "lines",
    "horizontal_strategy": "lines",
    "intersection_tolerance": 5,
    "join_tolerance": 5,
    "snap_tolerance": 3,
}


def _glue_for(column: str) -> str:
    return "" if any(k in column for k in GLUE_WITHOUT_SPACE) else " "


def _is_numeric_column(column: str) -> bool:
    return any(k in column for k in NUMERIC_HINTS)


def _to_number(value: str) -> Any:
    """'2,174.0' → 2174.0. 숫자가 아니면 원문 그대로."""
    text = value.replace(",", "").replace(" ", "")
    if not text:
        return ""
    if re.fullmatch(r"-?\d+(\.\d+)?", text):
        number = float(text)
        return int(number) if number.is_integer() and "." not in text else number
    return value


def _merge_rows(
    layout: TableLayout,
    body: Sequence[Sequence[Optional[str]]],
    page_no: int,
) -> Tuple[List[Tuple[str, Dict[str, Any]]], List[str]]:
    """여러 줄에 걸친 한 건을 한 행으로 합친다.

    인쇄된 순서대로 ('record'|'subtotal', 행) 목록을 돌려준다.
    """
    ordered: List[Tuple[str, Dict[str, Any]]] = []
    warnings: List[str] = []

    current: Optional[Dict[str, List[str]]] = None
    current_is_subtotal = False

    def flush() -> None:
        nonlocal current, current_is_subtotal
        if current is None:
            return
        row: Dict[str, Any] = {}
        for column in layout.columns:
            text = _glue_for(column).join(current[column]).strip()
            text = re.sub(r"\s{2,}", " ", text)
            row[column] = _to_number(text) if _is_numeric_column(column) else text
        row["페이지"] = page_no
        ordered.append(("subtotal" if current_is_subtotal else "record", row))
        current = None
        current_is_subtotal = False

    subtotal_keys = [s.replace(" ", "") for s in SUBTOTAL_LABELS]
    anchor_col = layout.anchor_col

    for raw in body:
        cells: List[str] = []
        for idx in layout.col_index:
            cells.append(_norm(raw[idx]) if idx < len(raw) else "")

        key = cells[layout.key_col].replace(" ", "")
        anchor = cells[anchor_col] if anchor_col is not None else ""
        is_subtotal = key in subtotal_keys

        # 번호는 셀 가운데에 인쇄되기 때문에, 소재지가 여러 줄인 항목에서는
        # 자료보다 한 줄 아래에 찍히기도 한다. 그때는 새 항목이 아니라
        # 진행 중인 항목의 번호로 본다.
        number_of_current = (
            current is not None
            and bool(key)
            and not is_subtotal
            and not anchor
            and not current[layout.columns[layout.key_col]]
        )

        # 번호가 있으면 새 항목. 번호가 없어도 지번이 다시 나오면 새 항목.
        starts_record = not number_of_current and (
            bool(key)
            or (
                anchor_col is not None
                and bool(anchor)
                and current is not None
                and bool(current[layout.columns[anchor_col]])
            )
        )

        if starts_record:
            flush()
            current = {c: [] for c in layout.columns}
            current_is_subtotal = is_subtotal
        elif current is None:
            # 페이지 첫 행이 앞 페이지 항목의 이어진 줄인 경우
            if any(cells):
                warnings.append(
                    f"{page_no}쪽: 번호 없는 행을 만나 앞 페이지 항목에 붙이지 못했습니다 → {cells}"
                )
            continue

        assert current is not None
        for column, value in zip(layout.columns, cells):
            if value:
                current[column].append(_flatten(value, _glue_for(column)))

    flush()
    return ordered, warnings


def extract(pdf_path: str) -> ExtractResult:
    """PDF에서 토지세목조서를 읽어 온다."""
    result = ExtractResult()
    layout: Optional[TableLayout] = None

    with pdfplumber.open(pdf_path) as pdf:
        result.total_pages = len(pdf.pages)
        for page_no, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables(TABLE_SETTINGS):
                page_layout = locate_layout(table)
                if page_layout is None:
                    continue  # 토지세목조서가 아닌 표(고시 본문의 도로구역 표 등)

                if layout is None:
                    layout = page_layout
                    result.columns = list(layout.columns)
                elif page_layout.signature() != layout.signature():
                    result.warnings.append(
                        f"{page_no}쪽: 앞 페이지와 표 머리글이 다릅니다 "
                        f"({' / '.join(page_layout.columns)}). 이 페이지 머리글 기준으로 읽습니다."
                    )

                body = table[page_layout.header_rows:]
                ordered, warns = _merge_rows(page_layout, body, page_no)
                result.sequence.extend(ordered)
                result.records.extend(row for kind, row in ordered if kind == "record")
                result.subtotals.extend(row for kind, row in ordered if kind == "subtotal")
                result.warnings.extend(warns)
                result.pages_with_table += 1

    if layout is None:
        raise ValueError(
            "이 PDF에서 토지세목조서 표를 찾지 못했습니다. "
            "스캔 이미지 PDF라면 먼저 OCR이 필요합니다."
        )

    # 소계 행에만 있고 자료 행에는 없는 열이 생길 수 있으니 열 목록을 맞춘다.
    for row in result.records + result.subtotals:
        for column in result.columns:
            row.setdefault(column, "")
    return result


# ---------------------------------------------------------------------------
# 검증 — 소계와 직접 합산값 비교
# ---------------------------------------------------------------------------


def _numeric_columns(columns: Sequence[str]) -> List[str]:
    return [c for c in columns if _is_numeric_column(c)]


def _group_column(columns: Sequence[str]) -> Optional[str]:
    for c in columns:
        if "소재지" in c or "시군" in c:
            return c
    return None


def verify(result: ExtractResult) -> List[Dict[str, Any]]:
    """소계가 나오는 순서대로 구간을 나누고, 그 구간 합계와 소계를 비교한다."""
    if not result.subtotals:
        return []

    num_cols = _numeric_columns(result.columns)
    group_col = _group_column(result.columns)
    ordered = result.sequence

    report: List[Dict[str, Any]] = []
    bucket: List[Dict[str, Any]] = []
    for kind, row in ordered:
        if kind == "record":
            bucket.append(row)
            continue
        entry: Dict[str, Any] = {
            "구간": bucket[0].get(group_col, "") if (bucket and group_col) else "",
            "건수": len(bucket),
            "소계 페이지": row["페이지"],
        }
        for col in num_cols:
            computed = sum(v for v in (r.get(col) for r in bucket) if isinstance(v, (int, float)))
            stated = row.get(col)
            entry[f"{col} 합산"] = round(computed, 4)
            entry[f"{col} 소계"] = stated
            if isinstance(stated, (int, float)):
                entry[f"{col} 차이"] = round(computed - stated, 4)
            else:
                entry[f"{col} 차이"] = ""
        report.append(entry)
        bucket = []

    if bucket:
        entry = {
            "구간": bucket[0].get(group_col, "") if group_col else "",
            "건수": len(bucket),
            "소계 페이지": "(소계 없음)",
        }
        for col in num_cols:
            computed = sum(v for v in (r.get(col) for r in bucket) if isinstance(v, (int, float)))
            entry[f"{col} 합산"] = round(computed, 4)
            entry[f"{col} 소계"] = ""
            entry[f"{col} 차이"] = ""
        report.append(entry)

    return report


# ---------------------------------------------------------------------------
# 엑셀 쓰기
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
BODY_FONT = Font(size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MAX_WIDTH = 42
MIN_WIDTH = 6


def _display_width(value: Any) -> int:
    """한글은 두 칸으로 세어 열 너비를 잡는다."""
    text = "" if value is None else str(value)
    return sum(2 if ord(ch) > 0x1100 else 1 for ch in text)


def _write_sheet(ws, columns: Sequence[str], rows: Sequence[Dict[str, Any]], number_format: str = "#,##0.0") -> None:
    ws.append(list(columns))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row in rows:
        ws.append([row.get(c, "") for c in columns])

    widths = [_display_width(c) for c in columns]
    numeric_flags = [_is_numeric_column(c) or c in ("건수", "페이지") for c in columns]

    for r in range(2, ws.max_row + 1):
        for c, column in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            if numeric_flags[c - 1] and isinstance(cell.value, (int, float)):
                cell.number_format = number_format if _is_numeric_column(column) else "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=False)
            widths[c - 1] = max(widths[c - 1], min(_display_width(cell.value), MAX_WIDTH * 2))

    for c, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(MIN_WIDTH, min(MAX_WIDTH, width / 1.6 + 2))

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{ws.max_row}"


def to_excel(result: ExtractResult, out_path: str, source_name: str = "") -> Dict[str, Any]:
    """추출 결과를 엑셀로 저장하고 요약 정보를 돌려준다."""
    checks = verify(result)

    wb = Workbook()

    columns = ["연번"] + list(result.columns) + ["페이지"]
    body = []
    for i, row in enumerate(result.records, start=1):
        item = {"연번": i}
        item.update(row)
        body.append(item)

    ws = wb.active
    ws.title = "토지세목조서"
    _write_sheet(ws, columns, body)

    if result.subtotals:
        ws_sub = wb.create_sheet("소계")
        _write_sheet(ws_sub, list(result.columns) + ["페이지"], result.subtotals)

    if checks:
        ws_chk = wb.create_sheet("검증")
        check_cols = list(checks[0].keys())
        _write_sheet(ws_chk, check_cols, checks)
        # 차이가 있는 칸은 빨갛게 — 고시문 소계와 어긋난 구간을 바로 알아보게 한다.
        red = Font(size=10, bold=True, color="C0392B")
        pink = PatternFill("solid", fgColor="FCE4E4")
        diff_cols = [i for i, c in enumerate(check_cols, start=1) if c.endswith("차이")]
        for r in range(2, ws_chk.max_row + 1):
            for c in diff_cols:
                cell = ws_chk.cell(row=r, column=c)
                if isinstance(cell.value, (int, float)) and abs(cell.value) > 0.05:
                    cell.font = red
                    for cc in range(1, len(check_cols) + 1):
                        ws_chk.cell(row=r, column=cc).fill = pink

    # 요약: 지목별 건수·면적
    num_cols = _numeric_columns(result.columns)
    jimok_col = next((c for c in result.columns if "지목" in c), None)
    summary_rows: List[Dict[str, Any]] = []
    if jimok_col:
        buckets: Dict[str, Dict[str, Any]] = {}
        for row in result.records:
            key = str(row.get(jimok_col, "")) or "(미기재)"
            entry = buckets.setdefault(key, {"지목": key, "건수": 0})
            entry["건수"] += 1
            for col in num_cols:
                value = row.get(col)
                if isinstance(value, (int, float)):
                    entry[col] = round(entry.get(col, 0) + value, 4)
        summary_rows = sorted(buckets.values(), key=lambda r: -r["건수"])
        total = {"지목": "합계", "건수": len(result.records)}
        for col in num_cols:
            total[col] = round(sum(r.get(col, 0) for r in summary_rows), 4)
        summary_rows.append(total)
        ws_sum = wb.create_sheet("지목별 요약")
        _write_sheet(ws_sum, ["지목", "건수"] + num_cols, summary_rows)

    wb.save(out_path)

    mismatches = [
        c for c in checks
        if any(isinstance(c.get(f"{col} 차이"), (int, float)) and abs(c[f"{col} 차이"]) > 0.05 for col in num_cols)
    ]
    return {
        "출력파일": out_path,
        "원본": source_name or "",
        "총 페이지": result.total_pages,
        "표가 있는 페이지": result.pages_with_table,
        "자료 건수": len(result.records),
        "소계 행": len(result.subtotals),
        "검증 구간": len(checks),
        "소계 불일치": len(mismatches),
        "불일치 구간": [str(m.get("구간", "")) for m in mismatches],
        "경고": result.warnings,
    }


def convert(pdf_path: str, out_path: Optional[str] = None) -> Dict[str, Any]:
    """PDF 한 개를 엑셀로 변환한다."""
    if out_path is None:
        out_path = os.path.splitext(pdf_path)[0] + ".xlsx"
    result = extract(pdf_path)
    return to_excel(result, out_path, source_name=os.path.basename(pdf_path))


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="수용할 또는 사용할 토지세목조서 PDF를 엑셀로 변환합니다.",
    )
    parser.add_argument("pdf", nargs="+", help="변환할 PDF 파일 (여러 개 가능, * 사용 가능)")
    parser.add_argument("-o", "--output", help="출력 파일 이름 (PDF 하나일 때만)")
    parser.add_argument("-d", "--outdir", help="출력 폴더 (여러 개 변환할 때)")
    args = parser.parse_args(argv)

    targets: List[str] = []
    for pattern in args.pdf:
        matched = sorted(glob.glob(pattern))
        targets.extend(matched or [pattern])

    if args.output and len(targets) > 1:
        parser.error("-o 는 PDF 가 하나일 때만 쓸 수 있습니다. 여러 개면 -d 를 쓰세요.")

    if args.outdir:
        os.makedirs(args.outdir, exist_ok=True)

    failures = 0
    for pdf_path in targets:
        if not os.path.exists(pdf_path):
            print(f"[건너뜀] 파일이 없습니다: {pdf_path}", file=sys.stderr)
            failures += 1
            continue

        if args.output:
            out_path = args.output
        else:
            name = os.path.splitext(os.path.basename(pdf_path))[0] + ".xlsx"
            out_path = os.path.join(args.outdir, name) if args.outdir else \
                os.path.join(os.path.dirname(pdf_path), name)

        try:
            info = convert(pdf_path, out_path)
        except Exception as exc:  # noqa: BLE001 - 사용자에게 그대로 보여준다
            print(f"[실패] {pdf_path}: {exc}", file=sys.stderr)
            failures += 1
            continue

        print(f"[완료] {pdf_path} → {info['출력파일']}")
        print(f"       자료 {info['자료 건수']:,}건 / 소계 {info['소계 행']}행 / "
              f"표 {info['표가 있는 페이지']}쪽 (전체 {info['총 페이지']}쪽)")
        if info["검증 구간"]:
            if info["소계 불일치"] == 0:
                print(f"       소계 검증: {info['검증 구간']}개 구간 모두 일치 (누락 없음)")
            else:
                print(f"       소계 검증: {info['검증 구간']}개 구간 중 "
                      f"{info['소계 불일치']}개가 고시문 소계와 다름 — 「검증」 시트 확인")
                for name in info["불일치 구간"]:
                    print(f"         · {name}")
        for warning in info["경고"][:10]:
            print(f"       ! {warning}")
        if len(info["경고"]) > 10:
            print(f"       ! 경고 {len(info['경고']) - 10}건 더 있음")

    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
